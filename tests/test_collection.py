import sys, unittest
from pathlib import Path
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from recruitment import registration,html_rows,table_jobs
from collector import read_document,recruitment_links,collect_document,Document
from eligibility import assess

class RegistrationTests(unittest.TestCase):
    def test_omitted_year_and_time(self):
        dates=registration('报名时间：2026年9月30日9:00起至10月16日17:00止')
        self.assertEqual(dates['opens'],'2026-09-30T09:00+08:00');self.assertEqual(dates['deadline'],'2026-10-16T17:00+08:00')
    def test_single_opening_is_not_a_deadline(self):
        self.assertIsNone(registration('报名时间：2026年10月8日开始受理')['deadline'])
    def test_immediate_start_requires_known_publication(self):
        self.assertIsNone(registration('报名时间：即日起至2026年10月21日')['opens'])
        self.assertEqual(registration('报名时间：即日起至2026年10月21日','2026-09-24')['opens'],'2026-09-24')
    def test_cross_year(self):
        self.assertEqual(registration('报名时间：2025年12月30日9:00至2026年1月12日17:00')['deadline'],'2026-01-12T17:00+08:00')
    def test_numeric_dates(self):
        self.assertEqual(registration('报名时间：2026/10/08至2026/10/13')['deadline'],'2026-10-13')
    def test_no_interview_date_leak(self):
        self.assertEqual(registration('报名时间：2026年10月8日至10月13日\n（二）面试时间\n2026年10月30日')['deadline'],'2026-10-13')

class CollectionTests(unittest.TestCase):
    def test_download_endpoint_attachment_is_attempted(self):
        import openpyxl,io
        book=openpyxl.Workbook();sheet=book.active
        sheet.append(['岗位名称','招聘人数','学历要求','专业要求']);sheet.append(['辅导员',2,'硕士','不限'])
        stream=io.BytesIO();book.save(stream)
        doc=Document('https://example.edu.cn/info/1/2.htm','2026年招聘公告','报名时间：2026年10月8日至10月13日','2026-09-20',[],[{'url':'https://example.edu.cn/download.jsp?id=1','title':'岗位表.xlsx'}],'h')
        jobs,_,errors=collect_document(doc,{'name':'测试院校','city':'广州市'},lambda _:stream.getvalue())
        self.assertEqual(len(jobs),1);self.assertFalse(errors)
    def test_attachment_failure_does_not_create_fictitious_jobs(self):
        doc=Document('https://example.edu.cn/info/1/2.htm','多岗位招聘公告','具体岗位见附件，报名时间：2026年10月8日至10月13日',None,[],[{'url':'https://example.edu.cn/download.jsp?id=1','title':'岗位表.xlsx'}],'h')
        jobs,_,errors=collect_document(doc,{'name':'测试院校'},lambda _:'<!doctype html><h1>验证码</h1>'.encode())
        self.assertFalse(jobs);self.assertTrue(errors)
    def test_html_merged_cells_and_private_school_metadata(self):
        raw='<table><tr><th>部门</th><th>岗位</th><th>人数</th><th>任职要求</th></tr><tr><td rowspan="2">学生处</td><td>辅导员</td><td>2</td><td>本科及以上学历，中共党员</td></tr><tr><td>宣传干事</td><td>1</td><td>硕士，传播学</td></tr></table>'.encode()
        doc=Document('https://example.edu.cn/info/1/2.htm','招聘公告','',None,[],[],'h')
        jobs=table_jobs(doc,{'name':'测试院校','nature':'民办','school_level':'专科'},html_rows(raw))
        self.assertEqual(len(jobs),2);self.assertEqual(jobs[1]['department'],'学生处');self.assertEqual(jobs[0]['nature'],'民办')
    def test_wp_article_and_page_url_are_discovered(self):
        raw='<title>招聘公告</title><div class="wp_articlecontent">岗位正文</div><a href="/2026/1001/c1a1/page.htm">教师招聘</a>'.encode()
        doc=read_document(raw,'https://example.edu.cn/')
        self.assertIn('岗位正文',doc.text);self.assertEqual(len(recruitment_links(doc)),1)
    def test_campus_fair_is_excluded(self):
        doc=read_document('<a href="/info/1/2.htm">校园招聘会通知</a>'.encode(),'https://example.edu.cn/')
        self.assertFalse(recruitment_links(doc))

class MatchingTests(unittest.TestCase):
    def test_teaching_is_not_implicitly_excluded(self):
        job={'degree':'硕士','majors':['传播学'],'category':'教师','title':'新媒体教师','conditions':'硕士传播学','evidence':{'degree':'硕士','major':'传播学'}}
        self.assertTrue(assess(job,{'degree':'硕士','major':'传播学','degree_origin':'境外'},'2026-10-09')['relevant'])
    def test_doctorate_remains_a_hard_limit(self):
        job={'degree':'博士','majors':['不限'],'category':'教师','title':'教师','evidence':{'degree':'博士','major':'专业不限'}}
        self.assertEqual(assess(job,{'degree':'硕士','major':'传播学'},'2026-10-09')['summary'],'不符合')
    def test_overseas_equivalence_remains_unconfirmed(self):
        job={'degree':'硕士','majors':['新闻传播学(A0503)'],'category':'教师','title':'教师','evidence':{'degree':'硕士','major':'新闻传播学(A0503)'}}
        result=assess(job,{'degree':'硕士','major':'传播学','degree_origin':'境外'},'2026-10-09')
        self.assertEqual(next(i['outcome'] for i in result['items'] if i['field']=='专业'),'需核实')
    def test_age_uses_announced_reference(self):
        job={'eligibility':{'age_text':'30周岁以下','age_as_of':'2026-10-01'}}
        result=assess(job,{'birth_date':'1990-01-01'},'2026-10-09')
        self.assertEqual(next(i['outcome'] for i in result['items'] if i['field']=='年龄'),'不符合')
class MergeTests(unittest.TestCase):
    def test_unchanged_auto_rows_are_reparsed(self):
        from scripts.collect_public import merge_job
        old={'id':'a','document_hash':'h','verification':'自动整理待核验','degree':None}
        new={**old,'degree':'博士'}
        self.assertEqual(merge_job(old,new)['degree'],'博士')
    def test_verified_conditions_are_retained(self):
        from scripts.collect_public import merge_job
        old={'id':'a','document_hash':'h','verification':'原文已核对','degree':'硕士'}
        new={**old,'document_hash':'new','degree':'博士'}
        result=merge_job(old,new)
        self.assertEqual(result['degree'],'硕士');self.assertEqual(result['verification'],'变更待复核')
    def test_duplicate_attachment_with_different_owner(self):
        from scripts.collect_public import deduplicate
        a={'id':'a','school':'测试','title':'宣传岗','source_url':'https://example.edu.cn/a','attachment_url':'https://example.edu.cn/download.jsp?owner=1&wbfileid=9','attachment_row':4}
        b={**a,'id':'b','source_url':'https://example.edu.cn/b','attachment_url':'https://example.edu.cn/download.jsp?owner=2&wbfileid=9'}
        rows=deduplicate([a,b]);self.assertEqual(len(rows),1);self.assertEqual(rows[0]['merged_ids'],['b'])
    def test_actual_spreadsheet_aliases_preserve_doctorate(self):
        import openpyxl,io
        from collector import parse_spreadsheet_jobs
        book=openpyxl.Workbook();sheet=book.active
        sheet.append(['序号','部门名称','岗位名称','招聘人数','学历','学位','专业名称及编号(须注明专业代码)','最低专业技术资格','与岗位有关的其它条件'])
        sheet.append([1,'宣传部','宣传人员',1,'研究生','博士','新闻传播学（A0503）','副高以上','35周岁以下；中共党员'])
        stream=io.BytesIO();book.save(stream)
        doc=Document('https://example.edu.cn/info/1/2.htm','招聘公告','',None,[],[],'h')
        job=parse_spreadsheet_jobs(doc,{'name':'测试'},stream.getvalue(),{'url':'https://example.edu.cn/file.xlsx'})[0]
        self.assertEqual(job['degree'],'博士');self.assertTrue(job['party_required']);self.assertIn('新闻传播学',job['majors'][0]);self.assertEqual(job['attachment_fields']['职称等级'],'副高以上')
    def test_affiliated_high_school_is_excluded(self):
        doc=read_document('<a href="/info/1/2.htm">深圳北理莫斯科大学附属实验中学教师招聘公告</a>'.encode(),'https://example.edu.cn/')
        self.assertFalse(recruitment_links(doc))
if __name__=='__main__':unittest.main()

